"""Generate precision Excel files — either from scratch (matching the lab
template cell-for-cell) or by populating the user's own uploaded template.

Port of the original TypeScript excel-export. Formulas (STDEV / AVERAGE / CV)
are written as live Excel formulas so the workbook recalculates on open.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ── Data containers ──────────────────────────────────────────────────────────

@dataclass
class AnalyteInfo:
    name: str
    levels: int
    unit: str = ""
    allowable_cv: float | None = None
    calibrator: str = ""
    reagent_lot: str = ""
    reagent_expiry: str = ""
    control_info: list[dict] = field(default_factory=list)


@dataclass
class WorklistInfo:
    name: str
    num_days: int
    reps_per_day: int
    intra_day_reps: int
    precision_type: str
    analytes: list[AnalyteInfo]
    metadata: dict = field(default_factory=dict)


@dataclass
class ResultRecord:
    analyte_name: str
    level: int
    day_number: int
    replication: int
    precision_type: str
    value: float | None
    date: str = ""
    unit: str = ""


# ── Styling ──────────────────────────────────────────────────────────────────

COL_TITLE = "FFD4E157"   # yellow-green title bar
COL_HEADER = "FF388E3C"  # green data header
COL_CTRL = "FFFFF9C4"    # pale yellow control header
COL_LABEL = "FFF5F5F5"
COL_WHITE = "FFFFFFFF"

_thin = Side(style="thin", color="FFBDBDBD")
_BORDER = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)


def _fill(argb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=argb)


def _set_title(cell, bg: str):
    cell.fill = _fill(bg)
    cell.font = Font(bold=True, size=11, color="FF1A1A1A")
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    cell.border = _BORDER


def _set_label(cell):
    cell.fill = _fill(COL_LABEL)
    cell.font = Font(bold=True, size=9)
    cell.alignment = Alignment(vertical="center")
    cell.border = _BORDER


def _set_value(cell):
    cell.font = Font(size=9)
    cell.alignment = Alignment(vertical="center")
    cell.border = _BORDER


def _set_header(cell):
    cell.fill = _fill(COL_HEADER)
    cell.font = Font(bold=True, size=9, color=COL_WHITE)
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    cell.border = _BORDER


def _set_data(cell, bold: bool = False):
    cell.font = Font(size=9, bold=bold)
    cell.alignment = Alignment(vertical="center", horizontal="center")
    cell.border = _BORDER
    cell.number_format = "0.####"


# ── Sheet builder — matches the lab template cell-for-cell ───────────────────

def _build_sheet(ws: Worksheet, analyte: AnalyteInfo, worklist: WorklistInfo,
                 results: list[ResultRecord], precision_type: str):
    meta = worklist.metadata or {}
    is_intra = precision_type == "intra"
    reps = max(worklist.intra_day_reps if is_intra else worklist.reps_per_day, 1)
    days = 1 if is_intra else max(worklist.num_days, 1)

    rep_start = 2
    rep_end = 1 + reps
    sd_col = 2 + reps
    mean_col = 3 + reps
    cv_col = 4 + reps
    allow_col = 5 + reps
    last_col = allow_col
    L = get_column_letter

    ws.column_dimensions["A"].width = 16
    for c in range(rep_start, rep_end + 1):
        ws.column_dimensions[L(c)].width = 13
    ws.column_dimensions[L(sd_col)].width = 10
    ws.column_dimensions[L(mean_col)].width = 10
    ws.column_dimensions[L(cv_col)].width = 10
    ws.column_dimensions[L(allow_col)].width = 13

    ws.row_dimensions[1].height = 6
    ws.merge_cells(f"A2:{L(last_col)}2")
    title = ws["A2"]
    title.value = " REPLICATION EXPERIMENT FOR PRECISION VERIFICATION"
    ws.row_dimensions[2].height = 22
    _set_title(title, COL_TITLE)
    for c in range(1, last_col + 1):
        ws[f"{L(c)}2"].border = _BORDER

    meta_rows = [
        ("Location", meta.get("location", "")),
        ("Instrument", meta.get("instrument", "")),
        ("HOD/COL", meta.get("hodcol", "")),
        ("Technician", meta.get("technician", "")),
        ("Test Name", analyte.name),
    ]
    r = 3
    for k, v in meta_rows:
        ws[f"A{r}"].value = k
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = v
        _set_value(ws[f"B{r}"])
        r += 1

    r = 9  # blank row 8, control block starts at 9

    for level in range(1, analyte.levels + 1):
        ctrl = (analyte.control_info[level - 1]
                if level - 1 < len(analyte.control_info) else {"lot": "", "limits": ""})

        ws[f"A{r}"].value = f" CONTROL {level}"
        _set_title(ws[f"A{r}"], COL_CTRL)
        ws[f"B{r}"].value = ctrl.get("lot", "")
        _set_value(ws[f"B{r}"])
        r += 1

        ws[f"A{r}"].value = "Control Limits"
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = ctrl.get("limits", "")
        _set_value(ws[f"B{r}"])
        r += 1

        ws[f"A{r}"].value = "Unit"
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = analyte.unit
        _set_value(ws[f"B{r}"])
        r += 1

        ws[f"A{r}"].value = "REAGENTS"
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = "LOT NO"
        _set_label(ws[f"B{r}"])
        ws[f"C{r}"].value = "EXPIRY DATE"
        _set_label(ws[f"C{r}"])
        r += 1

        ws[f"A{r}"].value = "Reagent Kit"
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = analyte.reagent_lot
        _set_value(ws[f"B{r}"])
        r += 1

        ws[f"A{r}"].value = "Controls"
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = ctrl.get("lot", "")
        _set_value(ws[f"B{r}"])
        ws[f"C{r}"].value = analyte.reagent_expiry
        _set_value(ws[f"C{r}"])
        r += 1

        ws[f"A{r}"].value = "Calibrator"
        _set_label(ws[f"A{r}"])
        ws[f"B{r}"].value = analyte.calibrator
        _set_value(ws[f"B{r}"])
        r += 1

        # Data header row
        ws[f"A{r}"].value = "DATE"
        _set_header(ws[f"A{r}"])
        for rep in range(1, reps + 1):
            c = rep_start + rep - 1
            ws[f"{L(c)}{r}"].value = f"REPLICATION-{rep}"
            _set_header(ws[f"{L(c)}{r}"])
        ws[f"{L(sd_col)}{r}"].value = "SD"
        _set_header(ws[f"{L(sd_col)}{r}"])
        ws[f"{L(mean_col)}{r}"].value = "MEAN"
        _set_header(ws[f"{L(mean_col)}{r}"])
        ws[f"{L(cv_col)}{r}"].value = "CV"
        _set_header(ws[f"{L(cv_col)}{r}"])
        ws[f"{L(allow_col)}{r}"].value = "Allowable CV"
        _set_header(ws[f"{L(allow_col)}{r}"])
        ws.row_dimensions[r].height = 20
        r += 1

        first_data_row = r

        for day in range(1, days + 1):
            day_results = [
                x for x in results
                if x.analyte_name == analyte.name
                and x.level == level
                and x.precision_type == precision_type
                and (x.day_number == 0 if is_intra else x.day_number == day)
            ]
            date_label = next((x.date for x in day_results if x.date), None)
            if date_label is None:
                date_label = "Run 1" if is_intra else f"Day {day}"
            ws[f"A{r}"].value = date_label
            _set_data(ws[f"A{r}"])

            for rep in range(1, reps + 1):
                c = rep_start + rep - 1
                found = next((x for x in day_results if x.replication == rep), None)
                cell = ws[f"{L(c)}{r}"]
                cell.value = found.value if found and found.value is not None else None
                _set_data(cell)

            rng = f"{L(rep_start)}{r}:{L(rep_end)}{r}"
            sd = ws[f"{L(sd_col)}{r}"]
            mean = ws[f"{L(mean_col)}{r}"]
            cv = ws[f"{L(cv_col)}{r}"]
            sd.value = f"=STDEV({rng})"
            mean.value = f"=AVERAGE({rng})"
            cv.value = f"=IF({L(mean_col)}{r}=0,0,{L(sd_col)}{r}/{L(mean_col)}{r}*100)"
            _set_data(sd)
            _set_data(mean)
            _set_data(cv)

            allow = ws[f"{L(allow_col)}{r}"]
            allow.value = analyte.allowable_cv
            _set_data(allow)
            r += 1

        last_data_row = r - 1

        ws[f"A{r}"].value = "Overall"
        _set_data(ws[f"A{r}"], True)
        for rep in range(1, reps + 1):
            _set_data(ws[f"{L(rep_start + rep - 1)}{r}"])
        block = f"{L(rep_start)}{first_data_row}:{L(rep_end)}{last_data_row}"
        o_sd = ws[f"{L(sd_col)}{r}"]
        o_mean = ws[f"{L(mean_col)}{r}"]
        o_cv = ws[f"{L(cv_col)}{r}"]
        o_sd.value = f"=STDEV({block})"
        o_mean.value = f"=AVERAGE({block})"
        o_cv.value = f"=IF({L(mean_col)}{r}=0,0,{L(sd_col)}{r}/{L(mean_col)}{r}*100)"
        _set_data(o_sd, True)
        _set_data(o_mean, True)
        _set_data(o_cv, True)
        _set_data(ws[f"{L(allow_col)}{r}"], True)
        r += 1


def build_workbook(analyte: AnalyteInfo, worklist: WorklistInfo,
                   results: list[ResultRecord]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet

    analyte_results = [r for r in results if r.analyte_name == analyte.name]
    has_inter = any(r.precision_type == "inter" for r in analyte_results)
    has_intra = any(r.precision_type == "intra" for r in analyte_results)

    if has_inter or (not has_inter and not has_intra):
        _build_sheet(wb.create_sheet("Inter-Day"), analyte, worklist, results, "inter")
    if has_intra:
        _build_sheet(wb.create_sheet("Intra-Day"), analyte, worklist, results, "intra")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Template ("dummy" Excel) support ─────────────────────────────────────────

@dataclass
class TemplateBlock:
    header_row: int
    rep_start_col: int
    rep_cols: int
    day_rows: list[int]
    unit_row: int | None


@dataclass
class ParsedTemplate:
    buffer: bytes
    sheet_name: str
    test_name_cell: tuple[int, int] | None
    blocks: list[TemplateBlock]
    max_reps: int
    max_days: int


def _cell_text(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _scan_sheet(ws: Worksheet):
    test_name_cell = None
    unit_rows: list[int] = []
    blocks: list[TemplateBlock] = []
    row_count = ws.max_row or 0

    for r in range(1, row_count + 1):
        a = _cell_text(ws.cell(row=r, column=1).value).lower()
        if not a:
            continue

        if a == "test name" and test_name_cell is None:
            test_name_cell = (r, 2)
        if a == "unit":
            unit_rows.append(r)

        if a == "date":
            rep_start_col = 2
            rep_cols = 0
            c = rep_start_col
            while re.search(r"replication", _cell_text(ws.cell(row=r, column=c).value), re.IGNORECASE):
                rep_cols += 1
                c += 1

            day_rows: list[int] = []
            dr = r + 1
            while dr <= row_count:
                dv = _cell_text(ws.cell(row=dr, column=1).value)
                if not dv:
                    break
                if re.match(r"^(control|reagents?|unit|calibrator|controls)\b", dv, re.IGNORECASE):
                    break
                day_rows.append(dr)
                dr += 1

            if rep_cols > 0 and day_rows:
                blocks.append(TemplateBlock(r, rep_start_col, rep_cols, day_rows, None))

    # Associate each block with the nearest "Unit" row above its header.
    for block in blocks:
        candidates = [u for u in unit_rows if u < block.header_row]
        block.unit_row = max(candidates) if candidates else None

    return test_name_cell, blocks


def parse_template(buffer: bytes) -> ParsedTemplate:
    wb = load_workbook(io.BytesIO(buffer))

    best = None
    for ws in wb.worksheets:
        test_name_cell, blocks = _scan_sheet(ws)
        if best is None or len(blocks) > len(best[2]):
            best = (ws, test_name_cell, blocks)

    ws = best[0] if best else wb.worksheets[0]
    test_name_cell = best[1] if best else None
    blocks = best[2] if best else []

    max_reps = max((b.rep_cols for b in blocks), default=0)
    max_days = max((len(b.day_rows) for b in blocks), default=0)

    return ParsedTemplate(buffer, ws.title, test_name_cell, blocks, max_reps, max_days)


def populate_template(tpl: ParsedTemplate, analyte: AnalyteInfo,
                      results: list[ResultRecord]) -> bytes:
    wb = load_workbook(io.BytesIO(tpl.buffer))
    ws = wb[tpl.sheet_name] if tpl.sheet_name in wb.sheetnames else wb.worksheets[0]

    if tpl.test_name_cell:
        ws.cell(row=tpl.test_name_cell[0], column=tpl.test_name_cell[1]).value = analyte.name

    analyte_results = [r for r in results if r.analyte_name == analyte.name]
    use_inter = any(r.precision_type == "inter" for r in analyte_results)

    def fill_row(row: int, block: TemplateBlock, rows: list[ResultRecord]):
        if not rows:
            return
        date_val = next((x.date for x in rows if x.date), None)
        if date_val:
            ws.cell(row=row, column=1).value = date_val
        for rep in range(1, block.rep_cols + 1):
            found = next((x for x in rows if x.replication == rep), None)
            if found and found.value is not None:
                ws.cell(row=row, column=block.rep_start_col + rep - 1).value = found.value

    for bi, block in enumerate(tpl.blocks):
        level = bi + 1
        if block.unit_row and analyte.unit:
            ws.cell(row=block.unit_row, column=2).value = analyte.unit

        if use_inter:
            for di, row in enumerate(block.day_rows):
                day = di + 1
                day_results = [
                    x for x in analyte_results
                    if x.precision_type == "inter" and x.level == level and x.day_number == day
                ]
                fill_row(row, block, day_results)
        else:
            if block.day_rows:
                run_results = [
                    x for x in analyte_results
                    if x.precision_type == "intra" and x.level == level
                ]
                fill_row(block.day_rows[0], block, run_results)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── ZIP helper ───────────────────────────────────────────────────────────────

def _safe(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "_", s, flags=re.IGNORECASE)


def build_zip(worklist: WorklistInfo, results: list[ResultRecord],
              template: ParsedTemplate | None = None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for analyte in worklist.analytes:
            data = (populate_template(template, analyte, results)
                    if template else build_workbook(analyte, worklist, results))
            zf.writestr(f"{_safe(analyte.name)}_Precision.xlsx", data)
    return buf.getvalue()


def analyte_filename(name: str) -> str:
    return f"{_safe(name)}_Precision.xlsx"
